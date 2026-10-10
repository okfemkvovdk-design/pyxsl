"""Remove fully blank rows from Excel workbooks and report duplicate rows."""

from __future__ import annotations

import argparse
import json
import sys
import warnings
import zipfile
from collections.abc import Sequence
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _row_values(worksheet: Worksheet, row_number: int) -> tuple[Any, ...]:
    values = tuple(
        None if _is_blank(cell.value) else cell.value for cell in worksheet[row_number]
    )
    while values and values[-1] is None:
        values = values[:-1]
    return values


def _duplicate_key(values: tuple[Any, ...]) -> tuple[Any, ...]:
    return tuple(
        ("boolean", value) if isinstance(value, bool) else value for value in values
    )


def _remove_blank_rows(worksheet: Worksheet) -> int:
    if not any(
        not _is_blank(cell.value) for row in worksheet.iter_rows() for cell in row
    ):
        return 0

    blank_rows = [
        row_number
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=worksheet.max_row),
            start=1,
        )
        if all(_is_blank(cell.value) for cell in row)
    ]
    for row_number in reversed(blank_rows):
        worksheet.delete_rows(row_number)
    return len(blank_rows)


def _find_duplicates(
    worksheet: Worksheet,
) -> list[tuple[int, int, int, tuple[Any, ...]]]:
    first_occurrences: dict[tuple[Any, ...], int] = {}
    occurrence_counts: dict[tuple[Any, ...], int] = {}
    duplicate_rows: list[tuple[int, int, tuple[Any, ...], tuple[Any, ...]]] = []

    for row_number in range(1, worksheet.max_row + 1):
        values = _row_values(worksheet, row_number)
        if not values:
            continue

        key = _duplicate_key(values)
        first_row = first_occurrences.get(key)
        if first_row is None:
            first_occurrences[key] = row_number
            occurrence_counts[key] = 1
        else:
            occurrence_counts[key] += 1
            duplicate_rows.append((row_number, first_row, key, values))

    return [
        (
            row_number,
            first_row,
            occurrence_counts[key],
            values,
        )
        for row_number, first_row, key, values in duplicate_rows
    ]


def _write_report(
    output_path: Path,
    duplicates: Sequence[tuple[str, int, int, int, tuple[Any, ...]]],
) -> None:
    report = Workbook()
    worksheet = report.active
    worksheet.title = "Duplicates"
    worksheet.append(
        [
            "Sheet",
            "Row number",
            "First occurrence row",
            "Occurrences in group",
            "Row values (JSON)",
        ]
    )

    for sheet_name, row_number, first_row, group_count, values in duplicates:
        row = worksheet.max_row + 1
        worksheet.cell(row=row, column=1, value=sheet_name)
        worksheet.cell(row=row, column=2, value=row_number)
        worksheet.cell(row=row, column=3, value=first_row)
        worksheet.cell(row=row, column=4, value=group_count)
        values_cell = worksheet.cell(
            row=row,
            column=5,
            value=json.dumps(values, ensure_ascii=False, default=str),
        )
        values_cell.data_type = "s"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    report.save(output_path)


def _numeric_values(worksheet: Worksheet, column: int) -> list[int | float]:
    return [
        cell.value
        for (cell,) in worksheet.iter_rows(
            min_row=2,
            min_col=column,
            max_col=column,
        )
        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
    ]


def _matches_filter(value: Any, expected: str) -> bool:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return value == float(expected)
        except ValueError:
            return False
    return str(value).strip().casefold() == expected.strip().casefold()


def _format_number(value: Decimal) -> str:
    result = format(value, "f")
    if "." in result:
        result = result.rstrip("0").rstrip(".")
    return result or "0"


def answer_question(worksheet: Worksheet, question: str) -> str:
    """Answer a simple question about a column whose name is in the first row."""
    parts = question.split(maxsplit=1)
    command = parts[0].casefold() if parts else ""
    argument = parts[1] if len(parts) == 2 else ""

    if command == "rows" and not argument:
        return str(
            sum(
                any(not _is_blank(cell.value) for cell in row)
                for row in worksheet.iter_rows(min_row=2)
            )
        )
    if command not in {"average", "count", "sum", "min", "max"}:
        raise ValueError("Use average, count, sum, min, max, rows, or quit.")
    if not argument.strip():
        raise ValueError(f"Add a column name after '{command}'.")

    column_name, has_value, match_value = argument.partition("=")
    column_name = column_name.strip()
    headers = [
        cell.value.strip()
        if isinstance(cell.value, str)
        else ""
        if cell.value is None
        else str(cell.value).strip()
        for cell in worksheet[1]
    ]
    matching_columns = [
        index
        for index, header in enumerate(headers, start=1)
        if header.casefold() == column_name.casefold()
    ]
    if not matching_columns:
        raise ValueError(
            f"Column '{column_name}' was not found. Available columns: "
            + ", ".join(header for header in headers if header)
        )
    if len(matching_columns) > 1:
        raise ValueError(f"Column '{column_name}' appears more than once.")

    column = matching_columns[0]
    cells = [
        cell.value
        for (cell,) in worksheet.iter_rows(
            min_row=2,
            min_col=column,
            max_col=column,
        )
    ]
    if command == "count":
        if has_value:
            expected = match_value.strip().casefold()
            if not expected:
                raise ValueError("Enter a value after '='.")
            count = sum(
                not _is_blank(value) and _matches_filter(value, match_value)
                for value in cells
            )
        else:
            count = sum(not _is_blank(value) for value in cells)
        return str(count)
    if has_value:
        raise ValueError(
            "Only count supports a value filter, such as count Status=Paid."
        )

    numbers = _numeric_values(worksheet, column)
    if not numbers:
        raise ValueError(f"Column '{column_name}' has no numeric values.")
    decimal_numbers = [Decimal(str(number)) for number in numbers]
    if command == "average":
        return _format_number(sum(decimal_numbers) / len(decimal_numbers))
    if command == "sum":
        return _format_number(sum(decimal_numbers))
    if command == "min":
        return _format_number(min(decimal_numbers))
    return _format_number(max(decimal_numbers))


def ask_questions(workbook_path: Path) -> None:
    workbook = load_workbook(
        workbook_path,
        data_only=True,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
    )
    print("Worksheets: " + ", ".join(workbook.sheetnames))
    try:
        while True:
            sheet_name = input("Choose a worksheet: ").strip()
            matching_sheet = next(
                (
                    name
                    for name in workbook.sheetnames
                    if name.casefold() == sheet_name.casefold()
                ),
                None,
            )
            if matching_sheet is not None:
                worksheet = workbook[matching_sheet]
                break
            print("Worksheet not found. Choose one of the listed names.")

        headers = [
            str(cell.value).strip()
            for cell in worksheet[1]
            if not _is_blank(cell.value)
        ]
        print("Columns: " + (", ".join(headers) if headers else "(none found)"))
        print(
            "Ask: average Column, count Column, count Column=value, "
            "sum Column, min Column, max Column, or rows. Type quit to exit."
        )
        while True:
            question = input("Data question: ").strip()
            if question.casefold() in {"quit", "exit"}:
                return
            try:
                print(answer_question(worksheet, question))
            except ValueError as error:
                print(f"Could not answer: {error}")
    except (EOFError, KeyboardInterrupt):
        print("\nLeaving question mode.")


def process_workbook(
    input_path: Path,
    cleaned_output: Path,
    duplicates_output: Path,
) -> tuple[int, int]:
    input_path = input_path.resolve()
    cleaned_output = cleaned_output.resolve()
    duplicates_output = duplicates_output.resolve()

    if input_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("Input must be an .xlsx or .xlsm workbook.")
    if not input_path.is_file():
        raise FileNotFoundError(f"Input workbook not found: {input_path}")
    if cleaned_output == input_path or duplicates_output == input_path:
        raise ValueError("Output paths must not overwrite the input workbook.")
    if cleaned_output == duplicates_output:
        raise ValueError("Cleaned workbook and duplicate report need different paths.")
    if cleaned_output.suffix.lower() != input_path.suffix.lower():
        raise ValueError("Cleaned workbook must use the same extension as the input.")
    if duplicates_output.suffix.lower() != ".xlsx":
        raise ValueError("Duplicate report must use the .xlsx extension.")

    workbook = load_workbook(
        input_path,
        keep_vba=input_path.suffix.lower() == ".xlsm",
    )
    has_formulas = any(
        cell.data_type == "f"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    )
    has_tables = any(worksheet.tables for worksheet in workbook.worksheets)
    blank_rows_removed = 0
    duplicates: list[tuple[str, int, int, int, tuple[Any, ...]]] = []

    for worksheet in workbook.worksheets:
        blank_rows_removed += _remove_blank_rows(worksheet)
        duplicates.extend(
            (worksheet.title, row_number, first_row, count, values)
            for row_number, first_row, count, values in _find_duplicates(worksheet)
        )

    if blank_rows_removed and (has_formulas or has_tables):
        warnings.warn(
            "Blank rows were removed from a workbook with formulas or tables. "
            "Review formula references and table ranges; openpyxl does not "
            "update all workbook references when rows are deleted.",
            UserWarning,
            stacklevel=2,
        )

    cleaned_output.parent.mkdir(parents=True, exist_ok=True)
    duplicates_output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(cleaned_output)
    _write_report(duplicates_output, duplicates)
    return blank_rows_removed, len(duplicates)


def _default_output_paths(input_path: Path) -> tuple[Path, Path]:
    return (
        input_path.with_name(f"{input_path.stem}_cleaned{input_path.suffix}"),
        input_path.with_name(f"{input_path.stem}_duplicates.xlsx"),
    )


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Create a cleaned workbook, report exact duplicate rows, and "
            "optionally answer questions using the input workbook."
        )
    )
    parser.add_argument("input", type=Path, help="Input .xlsx or .xlsm workbook")
    parser.add_argument(
        "--cleaned-output",
        type=Path,
        help="Path for the cleaned workbook (default: <name>_cleaned.<ext>)",
    )
    parser.add_argument(
        "--duplicates-output",
        type=Path,
        help="Path for the duplicate report (default: <name>_duplicates.xlsx)",
    )
    parser.add_argument(
        "--ask",
        action="store_true",
        help="Interactively ask simple questions about the input workbook",
    )
    args = parser.parse_args(argv)

    default_cleaned, default_duplicates = _default_output_paths(args.input)
    cleaned_output = args.cleaned_output or default_cleaned
    duplicates_output = args.duplicates_output or default_duplicates

    try:
        blank_rows_removed, duplicate_count = process_workbook(
            args.input,
            cleaned_output,
            duplicates_output,
        )
    except (
        OSError,
        ValueError,
        KeyError,
        InvalidFileException,
        zipfile.BadZipFile,
        ParseError,
    ) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1

    print(f"Cleaned workbook: {cleaned_output}")
    print(f"Duplicate report: {duplicates_output}")
    print(f"Fully blank rows removed: {blank_rows_removed}")
    print(f"Duplicate rows reported: {duplicate_count}")
    if args.ask:
        try:
            ask_questions(args.input)
        except (
            OSError,
            ValueError,
            KeyError,
            InvalidFileException,
            zipfile.BadZipFile,
            ParseError,
        ) as error:
            print(f"Error opening workbook for questions: {error}", file=sys.stderr)
            return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
