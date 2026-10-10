import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_data_control import answer_question, main, process_workbook


class ProcessWorkbookTests(unittest.TestCase):
    def test_answers_average_count_and_rows_questions(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Amount", "Status"])
        sheet.append([10, "Paid"])
        sheet.append([20, "paid"])
        sheet.append([None, "Due"])

        self.assertEqual(answer_question(sheet, "average Amount"), "15")
        self.assertEqual(answer_question(sheet, "count Status"), "3")
        self.assertEqual(answer_question(sheet, "count Status=paid"), "2")
        self.assertEqual(answer_question(sheet, "rows"), "3")
        self.assertEqual(answer_question(sheet, "AVERAGE\tAmount"), "15")

    def test_counts_numeric_values_as_numbers(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Amount"])
        sheet.append([1])
        sheet.append([1.0])
        sheet.append([2])

        self.assertEqual(answer_question(sheet, "count Amount=1.0"), "2")

    def test_numeric_answers_keep_decimal_precision(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Amount"])
        sheet.append([1234567.89])
        sheet.append([0])

        self.assertEqual(answer_question(sheet, "average Amount"), "617283.945")
        self.assertEqual(answer_question(sheet, "sum Amount"), "1234567.89")

    def test_numeric_answers_avoid_binary_float_artifacts(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Amount"])
        sheet.append([0.1])
        sheet.append([0.2])

        self.assertEqual(answer_question(sheet, "sum Amount"), "0.3")
        self.assertEqual(answer_question(sheet, "average Amount"), "0.15")

    def test_rejects_value_filters_for_numeric_operations(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Amount"])
        sheet.append([10])

        with self.assertRaisesRegex(ValueError, "Only count supports"):
            answer_question(sheet, "average Amount=10")

    def test_rejects_empty_count_filter(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Status"])
        sheet.append(["Paid"])

        with self.assertRaisesRegex(ValueError, "Enter a value"):
            answer_question(sheet, "count Status=")

    def test_ask_mode_runs_against_the_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            workbook = Workbook()
            workbook.active.title = "Orders"
            workbook.active.append(["Amount"])
            workbook.active.append([25])
            workbook.save(source)

            output = StringIO()
            with (
                patch(
                    "builtins.input", side_effect=["orders", "average Amount", "quit"]
                ),
                redirect_stdout(output),
            ):
                result = main([str(source), "--ask"])

            self.assertEqual(result, 0)
            self.assertIn("25", output.getvalue())
            self.assertTrue(source.with_name("source_cleaned.xlsx").exists())
            self.assertTrue(source.with_name("source_duplicates.xlsx").exists())

    def test_reports_unknown_column(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Amount"])

        with self.assertRaisesRegex(ValueError, "was not found"):
            answer_question(sheet, "average Missing")

    def test_removes_blank_rows_and_reports_duplicates_without_deleting_them(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            cleaned = root / "cleaned.xlsx"
            report = root / "duplicates.xlsx"

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Orders"
            sheet.append(["ID", "Amount"])
            sheet.append(["A-1", 25])
            sheet.append([" ", None])
            sheet.append(["A-1", 25])
            sheet.append(["A-1", 25])
            sheet.append(["flag", True])
            sheet.append(["flag", 1])
            other_sheet = workbook.create_sheet("Other")
            other_sheet.append(["A-1", 25])
            workbook.save(source)

            blank_count, duplicate_count = process_workbook(source, cleaned, report)

            self.assertEqual(blank_count, 1)
            self.assertEqual(duplicate_count, 2)

            cleaned_workbook = load_workbook(cleaned)
            self.assertEqual(
                list(cleaned_workbook["Orders"].values),
                [
                    ("ID", "Amount"),
                    ("A-1", 25),
                    ("A-1", 25),
                    ("A-1", 25),
                    ("flag", True),
                    ("flag", 1),
                ],
            )
            self.assertEqual(list(cleaned_workbook["Other"].values), [("A-1", 25)])

            report_sheet = load_workbook(report, data_only=True).active
            self.assertEqual(report_sheet.max_row, 3)
            rows = list(report_sheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual(
                [row[:4] for row in rows],
                [
                    ("Orders", 3, 2, 3),
                    ("Orders", 4, 2, 3),
                ],
            )

    def test_empty_worksheet_does_not_report_a_row_removed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "empty.xlsx"
            workbook = Workbook()
            workbook.save(source)

            blank_count, duplicate_count = process_workbook(
                source,
                root / "cleaned.xlsx",
                root / "duplicates.xlsx",
            )

            self.assertEqual(blank_count, 0)
            self.assertEqual(duplicate_count, 0)

    def test_rejects_output_path_that_would_overwrite_input(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            workbook = Workbook()
            workbook.save(source)

            with self.assertRaisesRegex(ValueError, "must not overwrite"):
                process_workbook(source, source, Path(directory) / "report.xlsx")

    def test_warns_when_deleting_blank_rows_from_a_formula_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            cleaned = root / "cleaned.xlsx"
            report = root / "duplicates.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Value"])
            sheet.append([None])
            sheet.append(["=1+1"])
            workbook.save(source)

            with self.assertWarnsRegex(UserWarning, "Review formula references"):
                process_workbook(source, cleaned, report)

            self.assertTrue(cleaned.is_file())

    def test_reports_invalid_workbook_without_traceback(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "invalid.xlsx"
            source.write_bytes(b"not an Excel workbook")
            errors = StringIO()

            with redirect_stdout(StringIO()), redirect_stderr(errors):
                result = main([str(source)])

            self.assertEqual(result, 1)
            self.assertIn("Error:", errors.getvalue())


if __name__ == "__main__":
    unittest.main()
